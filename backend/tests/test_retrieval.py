"""
Real ingestion + hybrid retrieval with the ONNX embedding model and ChromaDB.

Marked slow: the first run downloads the MiniLM model (~80 MB).
"""

import fitz
import pytest

from rag import retrieval
from rag.ingestion import delete_document_chunks, ingest_document

pytestmark = pytest.mark.slow


@pytest.fixture(scope="module")
def corpus(tmp_path_factory):
    import uuid

    import docx

    case_id = str(uuid.uuid4())
    folder = tmp_path_factory.mktemp("corpus")

    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "Parties: Acme Corp and Beta LLC. Governing law: Delaware.")
    pdf.new_page().insert_text((72, 72), "Confidentiality obligations survive for five years after termination.")
    pdf_path = folder / "nda.pdf"
    pdf.save(pdf_path)

    memo = docx.Document()
    memo.add_paragraph("Settlement memo: the indemnification cap is two million dollars.")
    table = memo.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Invoice"
    table.rows[0].cells[1].text = "Zephyr-77"
    docx_path = folder / "memo.docx"
    memo.save(docx_path)

    ids = {"pdf": str(uuid.uuid4()), "docx": str(uuid.uuid4())}
    pdf_result = ingest_document(str(pdf_path), case_id, ids["pdf"], filename="nda.pdf")
    docx_result = ingest_document(str(docx_path), case_id, ids["docx"], filename="memo.docx")
    return {"case_id": case_id, "ids": ids, "pdf": pdf_result, "docx": docx_result}


def test_ingestion_counts_pages_and_chunks(corpus):
    assert corpus["pdf"]["page_count"] == 2
    assert corpus["pdf"]["chunks_stored"] == 2, "each page is chunked separately"
    assert corpus["docx"]["chunks_stored"] >= 1


def test_semantic_query_finds_right_page(corpus):
    hits = retrieval.hybrid_retrieve("how long must information be kept secret", corpus["case_id"], top_k=3)
    assert "five years" in hits[0]["text"]
    assert (hits[0]["filename"], hits[0]["page_number"]) == ("nda.pdf", 2)


def test_keyword_query_finds_docx_table_cell(corpus):
    hits = retrieval.hybrid_retrieve("Zephyr-77", corpus["case_id"], top_k=3)
    assert hits[0]["filename"] == "memo.docx"
    assert "Zephyr-77" in hits[0]["text"]


def test_results_are_ranked_and_deduplicated(corpus):
    hits = retrieval.hybrid_retrieve("confidentiality", corpus["case_id"], top_k=10)
    scores = [h["score"] for h in hits]
    assert scores == sorted(scores, reverse=True)
    assert len({h["text"] for h in hits}) == len(hits)


def test_doc_filter(corpus):
    pdf_id = corpus["ids"]["pdf"]
    hits = retrieval.hybrid_retrieve("Zephyr-77 invoice", corpus["case_id"], top_k=5, doc_ids=[pdf_id])
    assert hits and all(h["doc_id"] == pdf_id for h in hits)


def test_unknown_case_returns_nothing():
    assert retrieval.hybrid_retrieve("anything", "no-such-case") == []


def test_bm25_cache_is_invalidated_on_delete(corpus):
    case_id = corpus["case_id"]
    retrieval.hybrid_retrieve("Zephyr-77", case_id)
    assert case_id in retrieval._cache

    assert delete_document_chunks(case_id, corpus["ids"]["docx"]) >= 1
    assert case_id not in retrieval._cache
    assert all("Zephyr" not in h["text"] for h in retrieval.hybrid_retrieve("Zephyr-77", case_id))
